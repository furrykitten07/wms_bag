import openpyxl
import pandas as pd
import json

print("=== 1. ANALYZING FIKRI.XLSX ===")
wb = openpyxl.load_workbook('data/FIKRI.xlsx', data_only=True)
print('Sheets:', wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    rows = list(ws.iter_rows(values_only=True))
    print(f"\n--- Sheet: {sname} (Total rows: {len(rows)}) ---")
    for idx, r in enumerate(rows[:15], 1):
        non_empty = [v for v in r if v is not None]
        if non_empty:
            print(f"Row {idx}: {r[:10]}")

# Read Sheet1 and CRITICAL
df_sheet1 = pd.read_excel('data/FIKRI.xlsx', sheet_name='Sheet1')
print("\nSheet1 Columns:", df_sheet1.columns.tolist())
print(df_sheet1.head(10))

df_critical = pd.read_excel('data/FIKRI.xlsx', sheet_name='CRITICAL')
print("\nCRITICAL Columns:", df_critical.columns.tolist())
print(df_critical)
