import openpyxl
import pandas as pd
import os

det_path = r"d:\Dev\wms-pt.-bag\data\det_spk_agustus 2026.xlsx"
tug_path = r"d:\Dev\wms-pt.-bag\data\TUG 5 agustus 2026.xlsx"

print("--- DET SPK AGUSTUS 2026 ---")
wb_det = openpyxl.load_workbook(det_path, data_only=True)
print("Sheet names:", wb_det.sheetnames)
for sname in wb_det.sheetnames:
    df = pd.read_excel(det_path, sheet_name=sname)
    print(f"\nSheet: {sname}, Shape: {df.shape}")
    print("Columns:", df.columns.tolist()[:15])
    print("Head:\n", df.head(3))

print("\n\n--- TUG 5 AGUSTUS 2026 ---")
wb_tug = openpyxl.load_workbook(tug_path, data_only=True)
print("Sheet names:", wb_tug.sheetnames)
for sname in wb_tug.sheetnames:
    df = pd.read_excel(tug_path, sheet_name=sname)
    print(f"\nSheet: {sname}, Shape: {df.shape}")
    print("Columns:", df.columns.tolist()[:15])
    print("Head:\n", df.head(3))
